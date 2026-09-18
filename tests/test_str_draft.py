"""
Unit and API tests for the STR draft feature (agent/str_draft.py,
agent/str_docx.py, and the /api/v1/agent/draft_str/{case_id} endpoints).

Verifies: per-sentence citation attachment, independent rejection of an
ungrounded sentence with a clear/actionable reason, structured field
extraction (accounts/transactions/amounts/date range), tier-aware missing
sections, DOCX generation with sourced/verifiable content, and that nothing
is persisted (identical input -> identical output, no server-side state).
"""

import os
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

os.environ.setdefault("AGENT_API_KEY", "test-key")

from agent.api import app  # noqa: E402
from agent.str_docx import build_str_docx, str_docx_filename  # noqa: E402
from agent.str_draft import build_str_draft  # noqa: E402

API_KEY_HEADERS = {"X-API-Key": os.environ["AGENT_API_KEY"]}
client = TestClient(app, headers=API_KEY_HEADERS)


def _ledger_trace_events() -> list[dict]:
    return [
        {
            "event_id": "EVT-101",
            "case_id": "CASE-LEDGER-002",
            "timestamp": "2026-09-16T14:15:00Z",
            "tool_called": "get_transaction",
            "tool_input": {"transaction_id": "TX-LEDGER-3011"},
            "tool_output_summary": (
                "Retrieved transaction TX-LEDGER-3011 (real_ledger): amount $12,500.00 (debit) "
                "at 2026-09-16T14:15:00Z."
            ),
            "narration_sentence": "Retrieved transaction TX-LEDGER-3011 for $12,500.00 processed on rail bank.xlsx.",
        },
        {
            "event_id": "EVT-102",
            "case_id": "CASE-LEDGER-002",
            "timestamp": "2026-09-16T14:20:00Z",
            "tool_called": "walk_graph",
            "tool_input": {"account_id": "ACC-1092", "tier": "real_ledger", "depth": 2},
            "tool_output_summary": (
                "Ledger analysis for account ACC-1092: evaluated 4 transactions. "
                "Initial balance $15,400.00, lowest balance $-3,200.00. Risk: 0.82."
            ),
            "narration_sentence": (
                "Account history analysis for ACC-1092 detected a severe anomaly: balance plummeted "
                "into negative overdraft ($-3,200.00) and high-velocity burst of 4 consecutive transfers."
            ),
        },
    ]


def _ledger_narrative_with_hallucination() -> str:
    grounded = " ".join(evt["narration_sentence"] for evt in _ledger_trace_events())
    return grounded + " This account is linked to an offshore shell company."


def test_grounded_sentences_get_citations():
    draft = build_str_draft(
        case_id="CASE-LEDGER-002",
        tier_origin="real_ledger",
        primary_transaction_id="TX-LEDGER-3011",
        risk_score=0.82,
        trace_events=_ledger_trace_events(),
        narrative=_ledger_narrative_with_hallucination(),
    )
    citations = draft["sentences_with_citations"]
    assert len(citations) == 2
    assert citations[0]["event_id"] == "EVT-101"
    assert citations[0]["tool_called"] == "get_transaction"
    assert citations[1]["event_id"] == "EVT-102"
    assert citations[1]["tool_called"] == "walk_graph"


def test_hallucinated_sentence_rejected_with_actionable_reason():
    draft = build_str_draft(
        case_id="CASE-LEDGER-002",
        tier_origin="real_ledger",
        primary_transaction_id="TX-LEDGER-3011",
        risk_score=0.82,
        trace_events=_ledger_trace_events(),
        narrative=_ledger_narrative_with_hallucination(),
    )
    notes = draft["validator_notes"]
    assert len(notes) == 1
    assert "offshore" in notes[0]["sentence"].lower()
    # Reason must name the specific unsupported claim, not just "rejected" —
    # this is what makes it actionable for an analyst reviewing the draft.
    # The sentence contains two flagged speculative terms ("offshore" and
    # "shell"); the checker reports the first one it finds — either is a
    # valid, actionable rejection reason.
    assert "offshore" in notes[0]["reason"].lower() or "shell" in notes[0]["reason"].lower()
    # The rejected sentence must never appear in the retained narrative.
    assert "offshore" not in draft["narrative"].lower()


def test_structured_fields_extracted_from_evidence():
    draft = build_str_draft(
        case_id="CASE-LEDGER-002",
        tier_origin="real_ledger",
        primary_transaction_id="TX-LEDGER-3011",
        risk_score=0.82,
        trace_events=_ledger_trace_events(),
        narrative=_ledger_narrative_with_hallucination(),
    )
    data = draft["exportable_data"]
    assert data["account_ids"] == ["ACC-1092"]
    assert data["transaction_ids"] == ["TX-LEDGER-3011"]
    assert "$12,500.00" in data["amounts_cited"]
    assert "$-3,200.00" in data["amounts_cited"]  # negative amounts must not be dropped
    assert data["date_range"] == {"start": "2026-09-16T14:15:00Z", "end": "2026-09-16T14:20:00Z"}
    assert any(r["code"] == "D" for r in data["reasons_for_suspicion"])
    assert any(r["code"] == "F" for r in data["reasons_for_suspicion"])  # risk_score 0.82 >= 0.75


def test_missing_sections_always_flag_absent_kyc_layer():
    draft = build_str_draft(
        case_id="CASE-LEDGER-002",
        tier_origin="real_ledger",
        primary_transaction_id="TX-LEDGER-3011",
        risk_score=0.82,
        trace_events=_ledger_trace_events(),
        narrative=_ledger_narrative_with_hallucination(),
    )
    gaps = " ".join(draft["exportable_data"]["missing_sections"])
    assert "Part 4" in gaps and "NOT AVAILABLE" in gaps
    assert "Part 5" in gaps


def test_card_tier_flagged_as_wrong_document_type():
    draft = build_str_draft(
        case_id="CASE-CARD-001",
        tier_origin="real_card",
        primary_transaction_id="TX-CARD-623",
        risk_score=0.93,
        trace_events=[
            {
                "event_id": "EVT-201",
                "case_id": "CASE-CARD-001",
                "timestamp": "2026-09-18T00:07:52Z",
                "tool_called": "get_transaction",
                "tool_input": {"transaction_id": "TX-CARD-623"},
                "tool_output_summary": "Retrieved transaction TX-CARD-623 for $529.00.",
                "narration_sentence": "Retrieved transaction TX-CARD-623 for $529.00 processed on rail creditcard.csv.",
            }
        ],
        narrative="Retrieved transaction TX-CARD-623 for $529.00 processed on rail creditcard.csv.",
    )
    gaps = " ".join(draft["exportable_data"]["missing_sections"])
    assert "not an AML behavioral pattern" in gaps
    assert draft["exportable_data"]["account_ids"] == []


def test_synthetic_tier_flagged_as_not_real():
    draft = build_str_draft(
        case_id="CASE-SYNTH-003",
        tier_origin="synthetic_network",
        primary_transaction_id="TX-SYNTH-5501",
        risk_score=0.94,
        trace_events=[
            {
                "event_id": "EVT-301",
                "case_id": "CASE-SYNTH-003",
                "timestamp": "2026-09-18T09:30:00Z",
                "tool_called": "walk_graph",
                "tool_input": {"account_id": "ACC-SYN-401", "tier": "synthetic_network", "depth": 3},
                "tool_output_summary": "Closed round-tripping loop across 3 accounts, $49,000 circulated.",
                "narration_sentence": "Graph traversal revealed a closed round-tripping topology circulating $49,000 across 3 intermediate accounts.",
            }
        ],
        narrative="Graph traversal revealed a closed round-tripping topology circulating $49,000 across 3 intermediate accounts.",
    )
    gaps = " ".join(draft["exportable_data"]["missing_sections"])
    assert "SYNTHETIC" in gaps
    assert any(r["code"] == "E" for r in draft["exportable_data"]["reasons_for_suspicion"])


def test_no_trace_events_produces_empty_but_valid_draft():
    draft = build_str_draft(
        case_id="CASE-EMPTY",
        tier_origin="real_ledger",
        primary_transaction_id="TX-LEDGER-0001",
        risk_score=0.5,
        trace_events=[],
        narrative="",
    )
    assert draft["sentences_with_citations"] == []
    assert draft["validator_notes"] == []
    assert draft["exportable_data"]["account_ids"] == []


def test_draft_is_deterministic_and_not_persisted():
    """Same input, called twice, must yield the same content (excluding the generated_at timestamp) — nothing is cached or stateful."""
    kwargs = dict(
        case_id="CASE-LEDGER-002",
        tier_origin="real_ledger",
        primary_transaction_id="TX-LEDGER-3011",
        risk_score=0.82,
        trace_events=_ledger_trace_events(),
        narrative=_ledger_narrative_with_hallucination(),
    )
    first = build_str_draft(**kwargs)
    second = build_str_draft(**kwargs)
    assert first["narrative"] == second["narrative"]
    assert first["sentences_with_citations"] == second["sentences_with_citations"]
    assert first["validator_notes"] == second["validator_notes"]


def test_docx_contains_sourced_sentences_and_review_notes():
    draft = build_str_draft(
        case_id="CASE-LEDGER-002",
        tier_origin="real_ledger",
        primary_transaction_id="TX-LEDGER-3011",
        risk_score=0.82,
        trace_events=_ledger_trace_events(),
        narrative=_ledger_narrative_with_hallucination(),
    )
    docx_bytes = build_str_docx(draft)
    doc = Document(BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    # Every retained sentence and its numbered citation marker is present.
    for idx, item in enumerate(draft["sentences_with_citations"], start=1):
        assert item["sentence"] in full_text
        assert f"[{idx}]" in full_text

    # The Sources section resolves each numbered marker to its event ID.
    assert "Event EVT-101" in full_text
    assert "Event EVT-102" in full_text

    # Review Notes section surfaces the rejected sentence and why.
    assert "REJECTED" in full_text
    assert "offshore" in full_text.lower()

    # Missing-KYC sections are visible in the exported document, not hidden.
    assert "NOT AVAILABLE" in full_text

    # Required footer disclaimer, verbatim.
    footer_text = doc.sections[0].footer.paragraphs[0].text
    assert footer_text == "Draft for analyst review. Generated by Verity. Not automatically filed."


def test_docx_filename_format():
    name = str_docx_filename("CASE-LEDGER-002")
    assert name.startswith("STR_CASE-LEDGER-002_")
    assert name.endswith(".docx")


def test_api_draft_str_endpoint_returns_citations_and_notes():
    payload = {
        "tier_origin": "real_ledger",
        "primary_transaction_id": "TX-LEDGER-3011",
        "risk_score": 0.82,
        "trace_events": _ledger_trace_events(),
        "narrative": _ledger_narrative_with_hallucination(),
    }
    response = client.post("/api/v1/agent/draft_str/CASE-LEDGER-002", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert len(data["sentences_with_citations"]) == 2
    assert len(data["validator_notes"]) == 1
    assert data["exportable_data"]["case_id"] == "CASE-LEDGER-002"


def test_api_draft_str_endpoint_requires_api_key():
    payload = {
        "tier_origin": "real_ledger",
        "primary_transaction_id": "TX-LEDGER-3011",
        "risk_score": 0.82,
        "trace_events": _ledger_trace_events(),
        "narrative": _ledger_narrative_with_hallucination(),
    }
    unauthenticated_client = TestClient(app)
    response = unauthenticated_client.post("/api/v1/agent/draft_str/CASE-LEDGER-002", json=payload)
    assert response.status_code == 401


def test_api_draft_str_docx_endpoint_returns_downloadable_file():
    payload = {
        "tier_origin": "real_ledger",
        "primary_transaction_id": "TX-LEDGER-3011",
        "risk_score": 0.82,
        "trace_events": _ledger_trace_events(),
        "narrative": _ledger_narrative_with_hallucination(),
    }
    response = client.post("/api/v1/agent/draft_str/CASE-LEDGER-002/docx", json=payload)
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "STR_CASE-LEDGER-002_" in response.headers["content-disposition"]

    doc = Document(BytesIO(response.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "TX-LEDGER-3011" in full_text
    assert "ACC-1092" in full_text
